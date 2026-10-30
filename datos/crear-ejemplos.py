#!/usr/bin/env python3
"""Genera los archivos de ejemplo (.docx, .pptx, .xlsx). Datos inventados."""
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

raiz = Path(__file__).resolve().parent
carpeta = raiz / "ejemplo"
carpeta.mkdir(parents=True, exist_ok=True)


# --------------------------------------------------------------------- .docx
doc = Document()

doc.add_heading("Informe de ejemplo", 0)
doc.add_paragraph(
    "Documento ficticio creado para probar el conversor. "
    "No contiene información de ninguna empresa real."
)

doc.add_heading("Objetivo", level=1)
p = doc.add_paragraph()
p.add_run("Verificar que la conversión conserva ")
p.add_run("la estructura").bold = True
p.add_run(", el ")
p.add_run("énfasis").italic = True
p.add_run(" y las tablas del documento original.")

doc.add_heading("Alcance", level=1)
doc.add_paragraph(
    "El documento cubre tres frentes de trabajo, en este orden de prioridad:",
)
for texto in [
    "Control de documentos por lotes",
    "Registro de cada conversión con fecha y origen",
    "Revisión manual de las tablas convertidas",
]:
    doc.add_paragraph(texto, style="List Bullet")

doc.add_heading("Resumen de datos", level=1)
doc.add_paragraph("Cifras agregadas del periodo, sin identificadores:")

tabla = doc.add_table(rows=4, cols=3)
tabla.style = "Table Grid"
filas = [
    ["Indicador", "Antes", "Después"],
    ["Documentos por lote", "12", "40"],
    ["Tiempo por documento", "38 min", "3 min"],
    ["Errores de transcripción", "6 por lote", "0"],
]
for f, fila in enumerate(filas):
    for c, valor in enumerate(fila):
        tabla.cell(f, c).text = valor

doc.add_heading("Notas", level=1)
doc.add_paragraph(
    "El texto de cierre es una nota al pie del responsable del proceso.",
    style="Quote",
)
doc.add_paragraph("Revisar las tablas con celdas combinadas.", style="List Number")
doc.add_paragraph("Confirmar la codificación de los acentos.", style="List Number")

doc.add_heading("Conclusión", level=1)
doc.add_paragraph(
    "La conversión a Markdown hace que el documento se pueda versionar con git, "
    "comparar entre versiones y revisar en el navegador."
)

ruta_docx = carpeta / "informe-ejemplo.docx"
doc.save(ruta_docx)


# --------------------------------------------------------------------- .pptx
prs = Presentation()

# Diapositiva 1: título
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = "Presentación de ejemplo"
slide.placeholders[1].text = "Documento ficticio para probar el conversor"

# Diapositiva 2: viñetas
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "Qué se automatiza"
cuerpo = slide.placeholders[1].text_frame
cuerpo.text = "Conversión por lotes de carpetas completas"
for linea in [
    "Registro de cada archivo procesado",
    "Validación de la jerarquía del resultado",
]:
    parrafo = cuerpo.add_paragraph()
    parrafo.text = linea

# Diapositiva 3: tabla
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "Resultados"
forma = slide.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(6), Inches(1.5))
tabla_pptx = forma.table
datos = [
    ["Formato", "Soportado"],
    ["Word", "Sí"],
    ["PowerPoint", "Sí"],
]
for f, fila in enumerate(datos):
    for c, valor in enumerate(fila):
        tabla_pptx.cell(f, c).text = valor

ruta_pptx = carpeta / "presentacion-ejemplo.pptx"
prs.save(ruta_pptx)


# --------------------------------------------------------------------- .xlsx
libro = Workbook()

hoja = libro.active
hoja.title = "Inventario"
hoja.append(["Producto", "Talla", "Stock", "Precio"])
hoja.append(["Camiseta local", "M", 24, 89900])
hoja.append(["Camiseta local", "L", 18, 89900])
hoja.append(["Sudadera", "M", 7, 159900])
hoja.append(["Medias", "Única", 120, 15000])

hoja2 = libro.create_sheet("Resumen")
hoja2.append(["Indicador", "Valor"])
hoja2.append(["Productos distintos", 3])
hoja2.append(["Unidades en stock", 169])

ruta_xlsx = carpeta / "inventario-ejemplo.xlsx"
libro.save(ruta_xlsx)

for ruta in (ruta_docx, ruta_pptx, ruta_xlsx):
    print(f"generado: {ruta.name} ({ruta.stat().st_size} bytes)")
