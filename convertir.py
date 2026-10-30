#!/usr/bin/env python3
"""
convertir.py — Convierte documentos de Office a Markdown con estructura jerárquica.

Formatos soportados: .docx, .pptx, .xlsx
Formatos detectados pero no soportados: .pdf, .doc

Uso:
    python convertir.py datos/ejemplo/
    python convertir.py datos/ejemplo/informe-ejemplo.docx
    python convertir.py archivo.pptx --salida salida/

Dependencias: python-docx, python-pptx, openpyxl (se importan solo si hacen falta).
"""

from __future__ import annotations

import argparse
import re
import sys
from datetime import date
from pathlib import Path

# Jerarquía del .docx: el título del documento queda como único H1 y los
# "Heading N" de Word bajan un nivel, para no producir varios H1 seguidos.
NIVEL_POR_ESTILO = {
    "Title": 1,
    "Heading 1": 2,
    "Heading 2": 3,
    "Heading 3": 4,
    "Heading 4": 5,
    "Heading 5": 6,
    "Heading 6": 6,
}

SOPORTADOS = {".docx", ".pptx", ".xlsx"}


# --------------------------------------------------------------------------
# Utilidades comunes
# --------------------------------------------------------------------------

def unir_bloques(bloques: list[list[str]]) -> str:
    """Une bloques de líneas separándolos con una línea en blanco.

    La separación no es cosmética: en Markdown, una lista pegada a una cita se
    absorbe dentro de la cita y deja de ser una lista.
    """
    limpios = [b for b in bloques if any(l.strip() for l in b)]
    return "\n\n".join("\n".join(b) for b in limpios).rstrip() + "\n"


def tabla_desde_filas(filas: list[list[str]]) -> list[str]:
    """Convierte filas de texto en una tabla de Markdown."""
    filas = [[c.replace("|", "\\|").replace("\n", " ").strip() for c in f] for f in filas]
    if not filas:
        return []

    ancho = max(len(f) for f in filas)
    filas = [f + [""] * (ancho - len(f)) for f in filas]

    lineas = ["| " + " | ".join(filas[0]) + " |"]
    lineas.append("|" + "|".join([" --- "] * ancho) + "|")
    for fila in filas[1:]:
        lineas.append("| " + " | ".join(fila) + " |")
    return lineas


def cabecera(ruta: Path) -> str:
    return f"<!-- convertido de {ruta.name} el {date.today().isoformat()} -->"


def importar(modulo: str, paquete: str):
    """Importa una librería solo cuando se necesita, con mensaje claro si falta."""
    try:
        return __import__(modulo, fromlist=["*"])
    except ImportError:
        raise SystemExit(
            f"Falta la librería '{paquete}' para leer este formato. "
            f"Instálala con: pip install {paquete}"
        )


# --------------------------------------------------------------------------
# .docx
# --------------------------------------------------------------------------

def _iter_bloques_docx(doc):
    """Recorre párrafos y tablas en el orden real del documento.

    Si se usara doc.paragraphs, todas las tablas quedarían al final y se perdería
    la posición que ocupan en el texto.
    """
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for hijo in doc.element.body.iterchildren():
        if hijo.tag == qn("w:p"):
            yield Paragraph(hijo, doc)
        elif hijo.tag == qn("w:tbl"):
            yield Table(hijo, doc)


def _run_a_markdown(run) -> str:
    texto = run.text
    if not texto.strip():
        return texto
    if run.bold and run.italic:
        return f"***{texto}***"
    if run.bold:
        return f"**{texto}**"
    if run.italic:
        return f"*{texto}*"
    return texto


def _nivel_de_lista(estilo: str) -> int:
    partes = estilo.split()
    if partes and partes[-1].isdigit():
        return int(partes[-1])
    return 1


def _clasificar_parrafo(parrafo) -> tuple[str, list[str]]:
    estilo = (parrafo.style.name or "").strip()
    texto = "".join(_run_a_markdown(r) for r in parrafo.runs).strip()

    if "w:drawing" in parrafo._p.xml:
        return "cita", ["> *(imagen no extraída)*"]

    if not texto:
        return "vacio", []

    if estilo in NIVEL_POR_ESTILO:
        return "titulo", [f"{'#' * NIVEL_POR_ESTILO[estilo]} {texto}"]

    if estilo.startswith("List Bullet"):
        return "lista_ul", ["  " * (_nivel_de_lista(estilo) - 1) + f"- {texto}"]

    if estilo.startswith("List Number"):
        return "lista_ol", ["  " * (_nivel_de_lista(estilo) - 1) + f"1. {texto}"]

    if estilo.startswith("Quote"):
        return "cita", [f"> {texto}"]

    return "texto", [texto]


def convertir_docx(ruta: Path) -> str:
    docx = importar("docx", "python-docx")
    doc = docx.Document(str(ruta))

    bloques: list[list[str]] = []
    lista_actual: list[str] = []
    tipo_lista = ""
    contador = 0

    def cerrar_lista():
        nonlocal lista_actual, tipo_lista, contador
        if lista_actual:
            bloques.append(lista_actual)
        lista_actual, tipo_lista, contador = [], "", 0

    for bloque in _iter_bloques_docx(doc):
        nombre = type(bloque).__name__

        if nombre == "Table":
            cerrar_lista()
            filas = [[c.text for c in fila.cells] for fila in bloque.rows]
            if filas:
                bloques.append(tabla_desde_filas(filas))
            continue

        tipo, lineas = _clasificar_parrafo(bloque)

        if tipo in ("lista_ul", "lista_ol"):
            if tipo != tipo_lista:
                cerrar_lista()
                tipo_lista = tipo
            contador += 1
            if tipo == "lista_ol":
                lista_actual.append(lineas[0].replace("1.", f"{contador}.", 1))
            else:
                lista_actual.append(lineas[0])
            continue

        cerrar_lista()
        if lineas:
            bloques.append(lineas)

    cerrar_lista()
    return cabecera(ruta) + "\n\n" + unir_bloques(bloques)


# --------------------------------------------------------------------------
# .pptx
# --------------------------------------------------------------------------

def convertir_pptx(ruta: Path) -> str:
    pptx = importar("pptx", "python-pptx")
    prs = pptx.Presentation(str(ruta))
    diapositivas = list(prs.slides)

    # El título del documento sale del título de la primera diapositiva; si no
    # tiene, se usa el nombre del archivo.
    titulo_documento = ruta.stem
    if diapositivas:
        primera = diapositivas[0].shapes.title
        if primera is not None and primera.has_text_frame:
            texto = primera.text_frame.text.strip()
            if texto:
                titulo_documento = texto

    bloques: list[list[str]] = [[f"# {titulo_documento}"]]

    for numero, slide in enumerate(diapositivas, start=1):
        titulo_shape = slide.shapes.title
        titulo = ""
        if titulo_shape is not None and titulo_shape.has_text_frame:
            titulo = titulo_shape.text_frame.text.strip()

        if numero == 1:
            # El título ya es el H1 del documento; repetirlo aquí sobraría.
            bloques.append(["## Diapositiva 1"])
        else:
            encabezado = f"## Diapositiva {numero}"
            if titulo:
                encabezado += f" — {titulo}"
            bloques.append([encabezado])

        for forma in slide.shapes:
            # Comparación por shape_id: python-pptx devuelve un objeto nuevo en
            # cada acceso a .title, así que comparar con `is` no funciona.
            if titulo_shape is not None and forma.shape_id == titulo_shape.shape_id:
                continue

            if getattr(forma, "has_table", False) and forma.has_table:
                filas = [[c.text for c in fila.cells] for fila in forma.table.rows]
                if filas:
                    bloques.append(tabla_desde_filas(filas))
                continue

            if getattr(forma, "has_text_frame", False) and forma.has_text_frame:
                items = []
                for parrafo in forma.text_frame.paragraphs:
                    texto = "".join(r.text for r in parrafo.runs).strip()
                    if texto:
                        items.append(f"- {texto}")
                if items:
                    bloques.append(items)

    return cabecera(ruta) + "\n\n" + unir_bloques(bloques)


# --------------------------------------------------------------------------
# .xlsx
# --------------------------------------------------------------------------

def convertir_xlsx(ruta: Path) -> str:
    openpyxl = importar("openpyxl", "openpyxl")
    libro = openpyxl.load_workbook(str(ruta), data_only=True)

    bloques: list[list[str]] = [[f"# {ruta.stem}"]]

    for hoja in libro.worksheets:
        filas: list[list[str]] = []
        for fila in hoja.iter_rows(values_only=True):
            celdas = ["" if v is None else str(v) for v in fila]
            if not any(c.strip() for c in celdas):
                continue
            filas.append(celdas)

        if not filas:
            continue

        bloques.append([f"## Hoja: {hoja.title}"])
        bloques.append(tabla_desde_filas(filas))

    return cabecera(ruta) + "\n\n" + unir_bloques(bloques)


# --------------------------------------------------------------------------
# Validación de estructura
# --------------------------------------------------------------------------

def validar_estructura(markdown: str) -> list[str]:
    """Revisa que la jerarquía de encabezados y la sangría sean consistentes.

    Es la comprobación que antes se hacía a mano antes de entregar los documentos
    al RAG: un salto de nivel o una sangría irregular hace que el documento se
    fragmente mal, y el recuperador devuelve pedazos sin contexto.
    """
    problemas: list[str] = []
    nivel_previo = 0
    hay_encabezado = False

    for numero, linea in enumerate(markdown.splitlines(), start=1):
        encabezado = re.match(r"^(#{1,6})\s+\S", linea)
        if encabezado:
            hay_encabezado = True
            nivel = len(encabezado.group(1))
            if nivel_previo == 0 and nivel != 1:
                problemas.append(
                    f"línea {numero}: el primer encabezado es H{nivel}, debería ser H1"
                )
            elif nivel_previo and nivel > nivel_previo + 1:
                problemas.append(
                    f"línea {numero}: salta de H{nivel_previo} a H{nivel}"
                )
            nivel_previo = nivel
            continue

        elemento_lista = re.match(r"^( +)[-*+]\s", linea)
        if elemento_lista:
            espacios = len(elemento_lista.group(1))
            if espacios % 2 != 0:
                problemas.append(
                    f"línea {numero}: sangría de {espacios} espacios "
                    "(se esperan múltiplos de 2)"
                )

    if not hay_encabezado:
        problemas.append("el documento no tiene ningún encabezado")

    return problemas


# --------------------------------------------------------------------------
# Interfaz
# --------------------------------------------------------------------------

CONVERSORES = {
    ".docx": convertir_docx,
    ".pptx": convertir_pptx,
    ".xlsx": convertir_xlsx,
}


def convertir(ruta: Path) -> str:
    conversor = CONVERSORES.get(ruta.suffix.lower())
    if conversor is None:
        raise SystemExit(f"Formato no soportado: {ruta.suffix}")
    return conversor(ruta)


def procesar(ruta: Path, carpeta_salida: Path | None) -> Path:
    markdown = convertir(ruta)
    destino = (carpeta_salida or ruta.parent) / (ruta.stem + ".md")
    destino.parent.mkdir(parents=True, exist_ok=True)
    destino.write_text(markdown, encoding="utf-8")

    problemas = validar_estructura(markdown)

    # Un archivo vacío o casi vacío es la falla más silenciosa de todas: el
    # proceso termina bien, el .md existe, y el RAG ingiere un documento sin
    # contenido. Se avisa explícitamente para que no pase inadvertido.
    cuerpo = markdown.split("\n\n", 1)[-1].strip()
    if len(cuerpo) < 20:
        problemas.insert(0, "la salida quedó vacía o casi vacía: el documento no produjo contenido")

    if problemas:
        for problema in problemas:
            print(f"  aviso: {problema}")
    else:
        print("  estructura OK")

    return destino


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Convierte .docx, .pptx y .xlsx a Markdown."
    )
    parser.add_argument("ruta", type=Path, help="archivo o carpeta")
    parser.add_argument(
        "--salida",
        type=Path,
        default=None,
        help="carpeta donde escribir los .md (por defecto, junto al original)",
    )
    args = parser.parse_args(argv)

    if not args.ruta.exists():
        print(f"No existe: {args.ruta}", file=sys.stderr)
        return 1

    if args.ruta.is_dir():
        todos = sorted(p for p in args.ruta.iterdir() if p.is_file())
        archivos = [p for p in todos if p.suffix.lower() in SOPORTADOS]
        # Solo se avisa de documentos que se esperaría convertir; cualquier otro
        # archivo de la carpeta se ignora en silencio.
        no_soportados = {".pdf", ".doc", ".ppt", ".xls", ".odt", ".rtf"}
        for archivo in todos:
            if archivo.suffix.lower() in no_soportados:
                print(f"omitido (formato no soportado): {archivo.name}")
    else:
        archivos = [args.ruta]

    if not archivos:
        print("No se encontraron archivos convertibles", file=sys.stderr)
        return 1

    for archivo in archivos:
        destino = procesar(archivo, args.salida)
        print(f"{archivo.name} -> {destino}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
